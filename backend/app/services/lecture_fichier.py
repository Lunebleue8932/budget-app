"""Ce qu'on sait d'un fichier AVANT d'en lire les colonnes : son format, son
encodage, son délimiteur, et pour un classeur, la feuille qui porte le relevé.

L'ENCODAGE. UTF-8 se reconnaît à coup sûr (un texte qui n'en est pas échoue
presque toujours à s'y décoder), avec ou sans BOM. Restent Windows-1252 et
ISO-8859-1, qui ne diffèrent que sur les octets 0x80 à 0x9F : des caractères
imprimables en Windows-1252 (« € », « ’ », « œ »), des codes de contrôle
invisibles en ISO-8859-1. Un fichier qui n'en contient aucun se lit pareil
dans les deux : on dit alors Windows-1252, l'encodage des exports bancaires
français sous Windows.

LA FEUILLE. Un classeur exporté porte parfois un onglet de synthèse, un de
légende, un de graphiques. La feuille PROPOSÉE est celle qui ressemble le plus
à un relevé — le plus de lignes portant à la fois une date et un nombre —, et
elle n'est que proposée : sans choix explicite, on lit la feuille active,
comme avant, et l'écran offre de changer.
"""
import codecs
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Optional

import openpyxl

from . import import_bancaire
from .lecture_montants import centimes

# Assez de lignes pour juger d'une feuille, assez peu pour qu'un classeur de
# dix onglets s'ouvre vite.
_LIGNES_PAR_FEUILLE = 300
# Les octets que Windows-1252 n'attribue à aucun caractère : leur présence
# exclut cet encodage.
_INDEFINIS_CP1252 = {0x81, 0x8D, 0x8F, 0x90, 0x9D}


def detecter_encodage(contenu: bytes) -> str:
    """« utf-8-sig », « utf-8 », « cp1252 » ou « latin-1 » (cf. le docstring
    du module pour la frontière entre les deux derniers)."""
    if contenu.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    try:
        contenu.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass
    if any(octet in _INDEFINIS_CP1252 for octet in contenu):
        return "latin-1"
    return "cp1252"


def _ressemble_a_une_operation(ligne) -> bool:
    """Une date et un nombre sur la même ligne : la forme d'une opération."""
    a_une_date = a_un_nombre = False
    for cellule in ligne:
        if cellule is None or cellule == "":
            continue
        if isinstance(cellule, (date, datetime)) or import_bancaire.parser_date(cellule) is not None:
            a_une_date = True
        elif centimes(cellule) is not None:
            a_un_nombre = True
    return a_une_date and a_un_nombre


@dataclass
class LectureFichier:
    lignes: list[tuple]
    format: str  # « csv » ou « xlsx »
    encodage: str = ""  # vide pour un classeur
    separateur: str = ""  # le délimiteur d'un CSV, « xlsx » pour un classeur
    # [{nom, operations}] : chaque feuille et le nombre de lignes qui y ont la
    # forme d'une opération. Vide pour un CSV.
    feuilles: list[dict] = field(default_factory=list)
    feuille: Optional[str] = None  # celle qui a été lue
    feuille_proposee: Optional[str] = None


def lire_fichier(
    contenu: bytes, delimiteur: Optional[str] = None, feuille: Optional[str] = None
) -> LectureFichier:
    """Les lignes du fichier, et ce qu'on a appris en les lisant. Les lignes
    sont EXACTEMENT celles que lit l'import (import_bancaire.
    _lire_toutes_les_lignes) : la détection ne doit pas décrire un autre
    fichier que celui qu'on importera."""
    if import_bancaire._est_xlsx(contenu):
        classeur = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True)
        feuilles = []
        for onglet in classeur.worksheets:
            lignes = list(onglet.iter_rows(values_only=True, max_row=_LIGNES_PAR_FEUILLE))
            feuilles.append({
                "nom": onglet.title,
                "operations": sum(_ressemble_a_une_operation(l) for l in lignes),
            })
        proposee = max(feuilles, key=lambda f: f["operations"])["nom"] if feuilles else None
        lue = classeur[feuille] if feuille and feuille in classeur.sheetnames else classeur.active
        return LectureFichier(
            lignes=list(lue.iter_rows(values_only=True)),
            format="xlsx",
            separateur="xlsx",
            feuilles=feuilles,
            feuille=lue.title,
            feuille_proposee=proposee,
        )
    texte = import_bancaire._decoder_texte(contenu)
    separateur = delimiteur or import_bancaire._detecter_delimiteur(
        "\n".join(texte.splitlines()[: import_bancaire._ECHANTILLON_DELIMITEUR])
    )
    return LectureFichier(
        lignes=import_bancaire._lire_lignes_csv(contenu, separateur),
        format="csv",
        encodage=detecter_encodage(contenu),
        separateur=separateur,
    )
