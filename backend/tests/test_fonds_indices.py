"""Les indices, le calque d'un ETF sur un indice, et le branchement d'un import sur des ETF détenus.

CE QUE CES TESTS PROTÈGENT :

  - UN INDICE NE SE DÉTIENT PAS : aucun achat ne le vise (route, crud), on ne transforme pas un
    titre détenu ni un titre qui a déjà porté des mouvements ;
  - LE CALQUE REMPLACE CE QUE L'UTILISATEUR A COCHÉ et rien d'autre : la composition d'un ETF reste
    si on ne la prend pas ; une propriété que l'indice ne porte pas est refusée plutôt que posée vide ;
  - UN IMPORT PEUT REMPLIR PLUSIEURS ETF DÉTENUS, chacun avec SES propriétés choisies, et `lire` dit
    ce que chacun porte déjà (c'est ce que l'écran montre avant d'écraser).
"""
import asyncio
import io
import json
from datetime import date

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app import crud, models, schemas
from app.constants import TYPE_COMPTE_PLACEMENT, AxeTitre, SensTitre
from app.services import calque_fonds

from .conftest import charger_module_extension, creer_compte

service = charger_module_extension("placements", "service_import_etf.py")
titres_routes = charger_module_extension("placements", "routeur_titres.py")
placements_routes = charger_module_extension("placements", "routeur_placements.py")


def _fichier(isin_fonds="IE000AAAAAA1", nom="Fonds Physique") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Actifs"
    for ligne in (
        [None, "Nom du fonds", nom],
        [None, "Code ISIN", isin_fonds],
        [None, "Devise de référence", "EUR"],
        [],
        [None, "Code ISIN", "Nom", "Classe d'actifs", "Devise", "Poids", "Secteur", "Pays"],
        [None, "US1912161007", "COCA-COLA", "EQUITY", "USD", 0.6, "Biens de consommation non cyclique", "États-Unis"],
        [None, "CH0038863350", "NESTLE", "EQUITY", "CHF", 0.4, "Biens de consommation non cyclique", "Suisse"],
    ):
        ws.append(ligne)
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def _monnaie(db):
    return crud.get_monnaies(db)[0]


def _etf(db, nom, isin=None):
    enveloppe = crud.get_types_titre(db, AxeTitre.enveloppe.value)
    etf_env = next((e for e in enveloppe if e.nom.lower().startswith("etf")), None) or crud.create_type_titre(
        db, "ETF", AxeTitre.enveloppe.value
    )
    return crud.create_titre(db, nom, _monnaie(db).id, 100.0, isin, etf_env.id)


def _compte(db):
    return db.query(models.Compte).filter(models.Compte.nom == "Compte-titres").first() or creer_compte(db, "Compte-titres", TYPE_COMPTE_PLACEMENT, solde_initial=10_000.0)


def _detenir(db, titre):
    """Un achat direct : on évite la route, qui exige un compte de placements."""
    compte = _compte(db)
    crud.create_operation_titre(
        db, compte_id=compte.id, titre=titre, sens=SensTitre.achat, quantite=2, prix_unitaire=10,
        date_operation=date(2026, 1, 5),
    )


# ---------- Les indices ----------


def test_un_indice_ne_s_achete_pas(db_session):
    indice = crud.create_titre(db_session, "MSCI World", _monnaie(db_session).id, 1000.0, est_indice=True)
    compte = _compte(db_session)
    with pytest.raises(ValueError, match="indice"):
        crud.create_operation_titre(
            db_session, compte_id=compte.id, titre=indice, sens=SensTitre.achat,
            quantite=1, prix_unitaire=1, date_operation=date(2026, 1, 5),
        )


def test_un_indice_se_cree_et_se_lit_comme_indice(db_session):
    payload = schemas.TitreCreate(nom="MSCI World", monnaie_id=_monnaie(db_session).id, valeur=3500, est_indice=True)
    lu = titres_routes.create_titre(payload, db=db_session)
    assert lu.est_indice is True and lu.valeur == 3500
    assert lu.detenu is False


def test_transformer_un_etf_non_detenu_en_indice(db_session):
    etf = _etf(db_session, "Fonds Physique", "IE000AAAAAA1")
    lu = titres_routes.update_titre(etf.id, schemas.TitreUpdate(est_indice=True), db=db_session)
    assert lu.est_indice is True
    # Et il peut redevenir un titre ordinaire.
    lu = titres_routes.update_titre(etf.id, schemas.TitreUpdate(est_indice=False), db=db_session)
    assert lu.est_indice is False


def test_on_ne_transforme_pas_un_titre_detenu_en_indice(db_session):
    etf = _etf(db_session, "Détenu", "IE000AAAAAA2")
    _detenir(db_session, etf)
    with pytest.raises(HTTPException) as erreur:
        titres_routes.update_titre(etf.id, schemas.TitreUpdate(est_indice=True), db=db_session)
    assert erreur.value.status_code == 409


def test_on_ne_transforme_pas_un_titre_qui_a_eu_des_mouvements(db_session):
    etf = _etf(db_session, "Vendu", "IE000AAAAAA3")
    _detenir(db_session, etf)
    compte = _compte(db_session)
    crud.create_operation_titre(
        db_session, compte_id=compte.id, titre=etf, sens=SensTitre.vente, quantite=2, prix_unitaire=12,
        date_operation=date(2026, 2, 5),
    )
    with pytest.raises(HTTPException) as erreur:
        titres_routes.update_titre(etf.id, schemas.TitreUpdate(est_indice=True), db=db_session)
    assert erreur.value.status_code == 409 and "mouvements" in erreur.value.detail


# ---------- Le calque ----------


def _indice_importe(db):
    resultat = service.confirmer(db, _fichier(), None, {"nb_lignes": 2})
    indice = crud.get_titre(db, resultat["etf_id"])
    titres_routes.update_titre(indice.id, schemas.TitreUpdate(est_indice=True), db=db)
    return crud.get_titre(db, indice.id)


def test_le_calque_prend_les_proprietes_cochees_et_laisse_les_autres(db_session):
    indice = _indice_importe(db_session)
    etf = _etf(db_session, "Synthétique", "IE000BBBBBB1")
    zone_a_moi = crud.create_type_titre(db_session, "Ma zone", AxeTitre.zone.value)
    crud.poser_repartition(db_session, etf, AxeTitre.zone.value, [(zone_a_moi.id, 100.0)])
    db_session.commit()

    lu = titres_routes.calquer_sur_indice(
        etf.id, schemas.CalqueTitre(indice_id=indice.id, proprietes=["composition", "secteur"]), db=db_session
    )
    assert set(lu.composition) == set(indice.composition) and lu.composition_couverture_pct == indice.composition_couverture_pct
    assert [e.type_titre_nom for e in lu.expositions if e.axe == "secteur"] == ["Consommation de base"]
    assert lu.expositions and any(e.axe == "secteur" and e.source.startswith("Calque sur") for e in lu.expositions)
    # La zone n'était pas cochée : celle de l'ETF reste.
    assert [e.type_titre_nom for e in lu.expositions if e.axe == "zone"] == ["Ma zone"]


def test_le_calque_exige_un_indice_et_une_propriete_que_l_indice_porte(db_session):
    indice = _indice_importe(db_session)
    etf = _etf(db_session, "Synthétique", "IE000BBBBBB1")
    pas_un_indice = _etf(db_session, "Autre", "IE000BBBBBB2")
    with pytest.raises(HTTPException) as erreur:
        titres_routes.calquer_sur_indice(
            etf.id, schemas.CalqueTitre(indice_id=pas_un_indice.id, proprietes=["composition"]), db=db_session
        )
    assert erreur.value.status_code == 404
    # L'indice importé n'a pas de classe d'actif propre : la demander est refusé.
    with pytest.raises(HTTPException) as erreur:
        titres_routes.calquer_sur_indice(
            etf.id, schemas.CalqueTitre(indice_id=indice.id, proprietes=["classe"]), db=db_session
        )
    assert erreur.value.status_code == 422
    with pytest.raises(HTTPException) as erreur:
        titres_routes.calquer_sur_indice(
            etf.id, schemas.CalqueTitre(indice_id=indice.id, proprietes=["couleur"]), db=db_session
        )
    assert erreur.value.status_code == 422


def test_un_titre_ne_se_cite_pas_lui_meme_dans_une_composition_posee(db_session):
    etf = _etf(db_session, "Synthétique", "IE000BBBBBB1")
    calque_fonds.poser_composition(etf, {"IE000BBBBBB1": 40.0, "US1912161007": 60.0}, 100.0)
    assert etf.composition == {"US1912161007": 60.0}


def test_resume_dit_ce_que_le_titre_porte(db_session):
    indice = _indice_importe(db_session)
    r = calque_fonds.resume(indice)
    assert r["composition"]["nb"] == 2 and r["composition"]["couverture_pct"] == pytest.approx(100.0)
    assert r["secteur"]["nb"] == 1 and r["classe"]["nb"] == 0


# ---------- L'import branché sur des ETF détenus ----------


def test_lire_propose_les_etf_a_brancher_detenus_d_abord_avec_ce_qu_ils_portent(db_session):
    detenu = _etf(db_session, "Synthétique détenu", "IE000BBBBBB1")
    _detenir(db_session, detenu)
    non_detenu = _etf(db_session, "Pas détenu", "IE000BBBBBB2")
    apercu = service.lire(db_session, _fichier())
    branchables = apercu["etf"]["branchables"]
    assert [b["id"] for b in branchables] == [detenu.id, non_detenu.id]
    assert [b["detenu"] for b in branchables] == [True, False]
    assert branchables[0]["donnees"]["composition"]["nb"] == 0
    # Un indice n'est jamais proposé : on y importe, on ne l'y branche pas.
    crud.create_titre(db_session, "MSCI World", _monnaie(db_session).id, 0, est_indice=True)
    assert len(service.lire(db_session, _fichier())["etf"]["branchables"]) == 2
    assert apercu["apport"] == {"composition": True, "secteur": True, "zone": True}


def test_confirmer_branche_les_etf_detenus_avec_leurs_proprietes(db_session):
    synthetique = _etf(db_session, "Synthétique", "IE000BBBBBB1")
    zone_a_moi = crud.create_type_titre(db_session, "Ma zone", AxeTitre.zone.value)
    crud.poser_repartition(db_session, synthetique, AxeTitre.zone.value, [(zone_a_moi.id, 100.0)])
    ancien = {"US0378331005": 100.0}
    synthetique.composition = ancien
    db_session.commit()
    autre = _etf(db_session, "Autre détenu", "IE000BBBBBB2")

    resultat = service.confirmer(
        db_session, _fichier(), None,
        {
            "nb_lignes": 2,
            "cibles": [
                # Le premier prend la composition et le secteur, garde sa zone.
                {"titre_id": synthetique.id, "proprietes": ["composition", "secteur"]},
                # Le second ne prend que la zone.
                {"titre_id": autre.id, "proprietes": ["zone"]},
            ],
        },
    )
    fonds = crud.get_titre(db_session, resultat["etf_id"])
    assert fonds.id not in (synthetique.id, autre.id) and set(fonds.composition) == {"US1912161007", "CH0038863350"}
    synthetique, autre = crud.get_titre(db_session, synthetique.id), crud.get_titre(db_session, autre.id)
    assert set(synthetique.composition) == set(fonds.composition)
    assert [e.etiquette.nom for e in synthetique.repartition(AxeTitre.secteur.value)] == ["Consommation de base"]
    assert [e.etiquette.nom for e in synthetique.repartition(AxeTitre.zone.value)] == ["Ma zone"]
    assert autre.composition is None and autre.repartition(AxeTitre.secteur.value) == []
    assert {e.etiquette.nom for e in autre.repartition(AxeTitre.zone.value)} >= {"États-Unis"}
    assert [b["nom"] for b in resultat["branches"]] == ["Synthétique", "Autre détenu"]


def test_ne_rien_cocher_sur_l_etf_du_fichier_le_laisse_intact(db_session):
    existant = _etf(db_session, "Fonds Physique", "IE000AAAAAA1")
    existant.composition = {"US0378331005": 100.0}
    existant.composition_couverture_pct = 100.0
    db_session.commit()
    service.confirmer(
        db_session, _fichier(), None,
        {"nb_lignes": 2, "etf": {"titre_id": existant.id, "proprietes": []}},
    )
    existant = crud.get_titre(db_session, existant.id)
    assert existant.composition == {"US0378331005": 100.0}
    assert existant.repartition(AxeTitre.secteur.value) == []


def test_on_ne_branche_pas_un_indice_ni_un_etf_disparu(db_session):
    indice = crud.create_titre(db_session, "MSCI World", _monnaie(db_session).id, 0, est_indice=True)
    with pytest.raises(ValueError, match="indice"):
        service.confirmer(db_session, _fichier(), None, {"nb_lignes": 2, "cibles": [{"titre_id": indice.id}]})
    db_session.rollback()
    with pytest.raises(ValueError, match="n'existe plus"):
        service.confirmer(db_session, _fichier(), None, {"nb_lignes": 2, "cibles": [{"titre_id": 99999}]})
    db_session.rollback()
    with pytest.raises(ValueError, match="inconnue"):
        service.confirmer(db_session, _fichier(), None, {"nb_lignes": 2, "etf": {"proprietes": ["couleur"], "nom": "X"}})


def test_donnees_du_titre_route(db_session):
    indice = _indice_importe(db_session)
    r = titres_routes.donnees_du_titre(indice.id, db=db_session)
    assert r["composition"]["nb"] == 2
    with pytest.raises(HTTPException):
        titres_routes.donnees_du_titre(99999, db=db_session)


def test_un_indice_n_a_pas_d_enveloppe_et_redevenir_ordinaire_la_rend(db_session):
    etf = _etf(db_session, "Fonds Physique", "IE000AAAAAA1")
    etf.composition = {"US0378331005": 100.0}
    db_session.commit()
    assert etf.type_titre_id is not None
    lu = titres_routes.update_titre(etf.id, schemas.TitreUpdate(est_indice=True), db=db_session)
    assert lu.est_indice is True and lu.type_titre_id is None
    # On ne lui en pose pas.
    with pytest.raises(HTTPException) as erreur:
        titres_routes.update_titre(etf.id, schemas.TitreUpdate(type_titre_id=1), db=db_session)
    assert erreur.value.status_code == 422
    with pytest.raises(HTTPException):
        titres_routes.create_titre(
            schemas.TitreCreate(nom="X", monnaie_id=_monnaie(db_session).id, est_indice=True, type_titre_id=1),
            db=db_session,
        )
    lu = titres_routes.update_titre(etf.id, schemas.TitreUpdate(est_indice=False), db=db_session)
    assert lu.est_indice is False and lu.type_titre_nom.lower().startswith("etf")


def test_le_calque_se_souvient_de_l_indice_et_se_repercute_a_sa_mise_a_jour(db_session):
    indice = _indice_importe(db_session)
    etf = _etf(db_session, "Synthétique", "IE000BBBBBB1")
    lu = titres_routes.calquer_sur_indice(
        etf.id, schemas.CalqueTitre(indice_id=indice.id, proprietes=["composition"]), db=db_session
    )
    assert lu.indice_id == indice.id and lu.indice_nom == indice.nom_affiche and lu.calque_proprietes == ["composition"]
    # L'indice est mis à jour par un nouvel import (sur la ligne existante) : l'ETF suit, sans refaire le lien.
    resultat = service.confirmer(
        db_session, _fichier(), None, {"nb_lignes": 1, "etf": {"titre_id": indice.id}}
    )
    assert resultat["etf_suivis"] == ["Synthétique"]
    assert set(crud.get_titre(db_session, etf.id).composition) == {"US1912161007"}
    # Délier : les données restent, le lien disparaît.
    lu = titres_routes.calquer_sur_indice(etf.id, schemas.CalqueTitre(indice_id=None), db=db_session)
    assert lu.indice_id is None and lu.composition


def test_supprimer_un_indice_delie_les_etf(db_session):
    indice = _indice_importe(db_session)
    etf = _etf(db_session, "Synthétique", "IE000BBBBBB1")
    titres_routes.calquer_sur_indice(
        etf.id, schemas.CalqueTitre(indice_id=indice.id, proprietes=["secteur"]), db=db_session
    )
    # Les composants sont cités par l'indice : on vide sa composition avant de le supprimer.
    indice.composition = None
    db_session.commit()
    crud.delete_titre(db_session, crud.get_titre(db_session, indice.id))
    assert crud.get_titre(db_session, etf.id).indice_id is None


def test_le_type_se_choisit_a_l_import(db_session):
    r = service.confirmer(db_session, _fichier(), None, {"nb_lignes": 2, "etf": {"nom": "Indice X", "isin": "", "indice": True}})
    indice = crud.get_titre(db_session, r["etf_id"])
    assert indice.est_indice and indice.type_titre_id is None and r["etf_est_indice"] is True
    env = crud.create_type_titre(db_session, "Fonds", AxeTitre.enveloppe.value)
    r = service.confirmer(
        db_session, _fichier("IE000CCCCCC1", "Autre"), None,
        {"nb_lignes": 2, "etf": {"nom": "Autre", "isin": "IE000CCCCCC1", "type_titre_id": env.id}},
    )
    assert crud.get_titre(db_session, r["etf_id"]).type_titre_id == env.id


def test_lire_propose_les_lignes_qui_ressemblent_au_fonds_du_fichier(db_session):
    proche = _etf(db_session, "AMUNDI FONDS PHYSIQUE ETF", "IE000ZZZZZZ1")
    _etf(db_session, "Tout autre chose", "IE000ZZZZZZ2")
    # Le nom d'affichage renommé à la main ne compte pas : on compare au nom d'origine.
    proche.nom_affichage = "Mon truc"
    db_session.commit()
    apercu = service.lire(db_session, _fichier(nom="Amundi Fonds Physique UCITS ETF"))
    suggestions = apercu["etf"]["suggestions"]
    assert suggestions and suggestions[0]["id"] == proche.id and suggestions[0]["score"] >= 0.6
    assert all(s["nom_courtier"] != "Tout autre chose" for s in suggestions)
    assert proche.id in [c["id"] for c in apercu["etf"]["candidats"]]


def test_le_meme_isin_est_une_ressemblance_parfaite(db_session):
    existant = _etf(db_session, "Nom différent", "IE000AAAAAA1")
    apercu = service.lire(db_session, _fichier())
    assert apercu["etf"]["suggestions"][0] == {
        "id": existant.id, "nom": "Nom différent", "nom_courtier": "Nom différent",
        "isin": "IE000AAAAAA1", "est_indice": False, "score": 1.0,
    }


# ---------- Le rapprochement flou des composants ----------

rapproche = charger_module_extension("placements", "rapprochement_titres.py")


def _fichier_sans_isin(*lignes, nom="Fonds iShares") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Holdings"
    ws.append(["Fund Name", nom])
    ws.append([])
    ws.append(["Ticker", "Name", "Sector", "Asset Class", "Weight (%)", "Location"])
    for ligne in lignes:
        ws.append(list(ligne))
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def test_les_libelles_de_bourse_se_ressemblent_malgre_les_suffixes():
    assert rapproche.score_noms("COCA-COLA CO/THE", "The Coca-Cola Company") >= 0.9
    assert rapproche.score_noms("PEPSICO INC", "PepsiCo, Inc.") >= 0.9
    assert rapproche.score_noms("NESTLE SA-REG", "Nestlé S.A.") >= 0.9
    assert rapproche.score_noms("COCA-COLA CO/THE", "PROCTER & GAMBLE CO/THE") < 0.5
    assert rapproche.score_noms("Bank of America", "Bank of China") < 0.8


def test_un_titre_connu_sous_un_autre_nom_est_rapproche_et_non_duplique(db_session):
    service.confirmer(db_session, _fichier(), None, {"nb_lignes": 2})
    coca = next(t for t in crud.get_titres(db_session) if t.code_isin == "US1912161007")
    avant = db_session.query(models.TitrePossede).count()
    fichier = _fichier_sans_isin(
        ("KO", "The Coca-Cola Company", "Consumer Staples", "Equity", 60.0, "United States"),
        ("NESN", "Nestlé S.A.", "Consumer Staples", "Equity", 40.0, "Switzerland"),
    )
    apercu = service.lire(db_session, fichier)
    proches = {r["nom"]: r for r in apercu["rapproches"]}
    assert proches["The Coca-Cola Company"]["titre_id"] == coca.id and proches["The Coca-Cola Company"]["accepte"]
    assert "même secteur" in proches["The Coca-Cola Company"]["raisons"]
    r = service.confirmer(db_session, fichier, None, {"nb_lignes": 2, "etf": {"nom": "iShares X", "isin": ""}})
    assert r["titres_rapproches"] == 2 and r["titres_crees"] == 0
    assert db_session.query(models.TitrePossede).count() == avant + 1  # seul le fonds est nouveau
    etf = crud.get_titre(db_session, r["etf_id"])
    assert f"US1912161007" in etf.composition


def test_un_rapprochement_refuse_cree_un_titre_et_un_secteur_contraire_le_decote(db_session):
    service.confirmer(db_session, _fichier(), None, {"nb_lignes": 2})
    fichier = _fichier_sans_isin(("KO", "The Coca-Cola Company", "Consumer Staples", "Equity", 100.0, "United States"))
    apercu = service.lire(db_session, fichier)
    ligne = apercu["rapproches"][0]["ligne"]
    r = service.confirmer(
        db_session, fichier, None,
        {"nb_lignes": 1, "etf": {"nom": "iShares X", "isin": ""}, "rapprochements": {str(ligne): False}},
    )
    assert r["titres_rapproches"] == 0 and r["titres_crees"] == 1
    # Le même nom dans un autre secteur et un autre pays pèse moins lourd.
    contraire = _fichier_sans_isin(("KO", "The Coca-Cola Company", "Financials", "Equity", 100.0, "Japan"))
    proche = service.lire(db_session, contraire)["rapproches"]
    assert not proche or proche[0]["score"] < apercu["rapproches"][0]["score"]


def test_deux_isin_differents_ne_se_rapprochent_jamais(db_session):
    service.confirmer(db_session, _fichier(), None, {"nb_lignes": 2})
    fichier = _fichier_sans_isin(("KO", "COCA-COLA CO/THE", "Consumer Staples", "Equity", 100.0, "United States"))
    # Avec un ISIN différent dans le fichier, ce serait un autre titre : simulé en posant l'ISIN sur la possession.
    p = rapproche.Rapprochement
    index = rapproche.IndexMots(crud.get_titres(db_session))
    coca = next(t for t in crud.get_titres(db_session) if t.code_isin == "US1912161007")
    possession = type("P", (), {"nom": "COCA-COLA CO/THE", "isin": "US0000000002"})()
    assert rapproche.chercher(index, possession, secteur_propose=None, zone_proposee=None, monnaie_du=None, cle=str.lower) is None
    possession.isin = "US1912161007"
    assert rapproche.chercher(index, possession, secteur_propose=None, zone_proposee=None, monnaie_du=None, cle=str.lower).titre is coca


def _fichier_ishares_multi_blocs() -> bytes:
    """La structure du fichier iShares : une feuille « Exposure Breakdowns » qui empile des blocs de
    répartition sous des titres, sans ligne d'en-tête, et le nom / l'ISIN du fonds dans d'autres feuilles."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Fund Header").append(["iShares MSCI World Health Care Sector Advanced UCITS ETF"])
    cles = wb.create_sheet("Key Facts")
    cles.append(["ISIN", "IE0009QS7W62", None])
    cles.append(["Base Currency", "USD", None])
    ex = wb.create_sheet("Exposure Breakdowns")
    for ligne in (
        ["Sector", None, None], ["as of", "15/Oct/2026", None], ["Type", "Fund", None],
        ["Health Care", "99.52", None], ["Cash and/or Derivatives", "0.48", None], [None, None, None], [None, None, None],
        ["Geography/Locations", None, None], ["as of", "15/Oct/2026", None], ["Type", "Fund", None],
        ["United States", "70.79", "US"], ["Switzerland", "12.42", "CH"], ["Other", "1.29", "OTHR"],
        ["Cash and/or Derivatives", "0.48", "CASHD_C"], [None, None, None],
        ["Allocations are subject to change.", None, None],
    ):
        ex.append(ligne)
    h = wb.create_sheet("Holdings")
    h.append(["All", None, None, None, None, None, None, None, None])
    h.append(["as of", "15/Oct/2026"])
    h.append(["Issuer Ticker", "Name", "Sector", "Asset Class", "Market Value", "Weight (%)", "Notional Value", "Nominal", "Market Currency"])
    h.append(["LLY", "ELI LILLY", "Health Care", "Equity", "USD 1", "60", "1", "1", "USD"])
    h.append(["NOVN", "NOVARTIS AG", "Health Care", "Equity", "USD 1", "40", "1", "1", "CHF"])
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def test_la_feuille_exposure_breakdowns_a_plusieurs_blocs_est_reconnue():
    fichier = charger_module_extension("placements", "composition_etf.py").lire(_fichier_ishares_multi_blocs())
    types = {f.nom.split(" — ")[-1]: f.type for f in fichier.feuilles}
    assert types["Sector"] == "secteurs" and types["Geography/Locations"] == "zones"
    # Liquidités et « Other » ne sont pas des parts classables.
    assert fichier.secteurs == [("Health Care", pytest.approx(99.52))]
    assert [lib for lib, _ in fichier.zones] == ["United States", "Switzerland"]
    assert fichier.fonds_nom.startswith("iShares MSCI World Health Care") and fichier.fonds_isin == "IE0009QS7W62"
    assert fichier.fonds_devise == "USD"


def test_la_repartition_du_fonds_vient_de_la_feuille_de_blocs(db_session):
    apercu = service.lire(db_session, _fichier_ishares_multi_blocs())
    assert apercu["a_une_repartition_secteurs"] and apercu["a_une_repartition_zones"]
    r = service.confirmer(db_session, _fichier_ishares_multi_blocs(), None, {"nb_lignes": 2})
    etf = crud.get_titre(db_session, r["etf_id"])
    assert [(e.etiquette.nom, e.poids_pct) for e in etf.repartition(AxeTitre.secteur.value)] == [("Santé", pytest.approx(99.52))]
    assert {e.etiquette.nom for e in etf.repartition(AxeTitre.zone.value)} >= {"États-Unis"}
